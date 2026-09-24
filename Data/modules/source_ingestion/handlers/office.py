"""Office document handlers — DOCX / XLSX / PPTX (optional deps)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from Data.modules.common.hashing import sha256_text

from ..settings import SourceIngestionSettings
from ..types import (
    ContentRef,
    DetectionResult,
    MemberOutcome,
    NormalizedArtifact,
    PARSER_VERSION,
    SourceKind,
)
from .base import HandlerCapabilities


class OfficeHandler:
    capabilities = HandlerCapabilities(
        handler_id="office",
        kinds=frozenset({SourceKind.OFFICE}),
        extensions=frozenset({".docx", ".xlsx", ".pptx"}),
        requires_path=True,
    )

    def can_handle(self, detection: DetectionResult, *, filename: str) -> bool:
        if detection.kind == SourceKind.OFFICE:
            return True
        return detection.extension.lower() in self.capabilities.extensions

    def inspect(self, path: Path, *, detection: DetectionResult, settings: SourceIngestionSettings) -> dict[str, Any]:
        return {"handler": "office", "extension": detection.extension}

    def ingest(
        self,
        path: Path,
        *,
        detection: DetectionResult,
        relative_path: str,
        settings: SourceIngestionSettings,
        staging_root: Path | None = None,
    ) -> NormalizedArtifact:
        ext = detection.extension.lower() or Path(relative_path).suffix.lower()
        try:
            if ext == ".docx":
                return self._docx(path, relative_path=relative_path)
            if ext == ".xlsx":
                return self._xlsx(path, relative_path=relative_path, settings=settings)
            if ext == ".pptx":
                return self._pptx(path, relative_path=relative_path)
        except ImportError as exc:
            return NormalizedArtifact(
                source_kind=SourceKind.OFFICE,
                title=Path(relative_path).name,
                relative_path=relative_path,
                mime_type=detection.mime_type,
                parser="office",
                parser_version=PARSER_VERSION,
                content_hash=sha256_text(str(path)),
                content=ContentRef(text=""),
                outcome=MemberOutcome.FAILED,
                error_code="OFFICE_DEPENDENCY_MISSING",
                skip_reason=str(exc)[:200],
                retryable=True,
                unsupported_features=[ext.lstrip(".")],
            )
        except Exception as exc:  # noqa: BLE001
            return NormalizedArtifact(
                source_kind=SourceKind.OFFICE,
                title=Path(relative_path).name,
                relative_path=relative_path,
                mime_type=detection.mime_type,
                parser="office",
                parser_version=PARSER_VERSION,
                content_hash=sha256_text(str(path)),
                content=ContentRef(text=""),
                outcome=MemberOutcome.FAILED,
                error_code="SOURCE_PARSE_FAILED",
                skip_reason=str(exc)[:200],
                retryable=True,
            )
        return NormalizedArtifact(
            source_kind=SourceKind.OFFICE,
            title=Path(relative_path).name,
            relative_path=relative_path,
            mime_type=detection.mime_type,
            parser="office",
            parser_version=PARSER_VERSION,
            content_hash=sha256_text(""),
            content=ContentRef(text=""),
            outcome=MemberOutcome.SKIPPED,
            skip_reason="unsupported_office_type",
        )

    def _docx(self, path: Path, *, relative_path: str) -> NormalizedArtifact:
        from docx import Document  # type: ignore[import-not-found]

        doc = Document(str(path))
        parts: list[str] = []
        para_meta: list[dict[str, Any]] = []
        for idx, para in enumerate(doc.paragraphs, start=1):
            text = para.text or ""
            if text.strip():
                parts.append(text)
                para_meta.append({"paragraph": idx, "chars": len(text)})
        for t_idx, table in enumerate(doc.tables, start=1):
            for r_idx, row in enumerate(table.rows, start=1):
                for c_idx, cell in enumerate(row.cells, start=1):
                    text = cell.text or ""
                    if text.strip():
                        parts.append(text)
                        para_meta.append({"table": t_idx, "row": r_idx, "col": c_idx, "chars": len(text)})
        body = "\n".join(parts)
        return NormalizedArtifact(
            source_kind=SourceKind.OFFICE,
            title=Path(relative_path).name,
            relative_path=relative_path,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            parser="docx",
            parser_version=PARSER_VERSION,
            content_hash=sha256_text(body),
            content=ContentRef(text=body),
            structured_metadata={"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)},
            provenance={"relative_path": relative_path, "locations": para_meta[:500]},
            outcome=MemberOutcome.SUCCESS,
        )

    def _xlsx(
        self,
        path: Path,
        *,
        relative_path: str,
        settings: SourceIngestionSettings,
    ) -> NormalizedArtifact:
        from openpyxl import load_workbook  # type: ignore[import-not-found]

        # Huge workbook → route to dataset when oversized.
        size = path.stat().st_size
        if size >= settings.dataset_route_min_bytes:
            return NormalizedArtifact(
                source_kind=SourceKind.DATASET,
                title=Path(relative_path).name,
                relative_path=relative_path,
                mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                parser="xlsx_dataset_route",
                parser_version=PARSER_VERSION,
                content_hash=sha256_text(f"xlsx:{path.stat().st_size}"),
                content=ContentRef(text=""),
                outcome=MemberOutcome.ROUTED,
                route_target="dataset",
                skip_reason="large_workbook_routed_to_dataset",
                provenance={"relative_path": relative_path, "size_bytes": size},
            )

        wb = load_workbook(str(path), read_only=True, data_only=True)
        parts: list[str] = []
        sheet_meta: list[dict[str, Any]] = []
        try:
            for sheet in wb.worksheets:
                sheet_meta.append({"sheet": sheet.title})
                parts.append(f"[sheet {sheet.title}]")
                row_count = 0
                for row in sheet.iter_rows(values_only=True):
                    row_count += 1
                    if row_count > 5000:
                        parts.append(f"... truncated after 5000 rows on sheet {sheet.title}")
                        # Honest partial — mark warning, do not claim full ingest of huge sheet as one blob.
                        break
                    cells = ["" if c is None else str(c) for c in row]
                    if any(cells):
                        parts.append("\t".join(cells))
                sheet_meta[-1]["rows_read"] = row_count
        finally:
            wb.close()
        body = "\n".join(parts)
        return NormalizedArtifact(
            source_kind=SourceKind.OFFICE,
            title=Path(relative_path).name,
            relative_path=relative_path,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            parser="xlsx",
            parser_version=PARSER_VERSION,
            content_hash=sha256_text(body),
            content=ContentRef(text=body),
            structured_metadata={"sheets": sheet_meta},
            provenance={"relative_path": relative_path, "sheets": sheet_meta},
            warnings=["xlsx_row_cap_5000_per_sheet"] if any(s.get("rows_read", 0) > 5000 for s in sheet_meta) else [],
            outcome=MemberOutcome.SUCCESS,
        )

    def _pptx(self, path: Path, *, relative_path: str) -> NormalizedArtifact:
        from pptx import Presentation  # type: ignore[import-not-found]

        prs = Presentation(str(path))
        parts: list[str] = []
        slides_meta: list[dict[str, Any]] = []
        for s_idx, slide in enumerate(prs.slides, start=1):
            slide_texts: list[str] = []
            for sh_idx, shape in enumerate(slide.shapes, start=1):
                if hasattr(shape, "text") and shape.text:
                    slide_texts.append(shape.text)
                    parts.append(f"[slide {s_idx} shape {sh_idx}]\n{shape.text}")
            slides_meta.append({"slide": s_idx, "shapes_with_text": len(slide_texts)})
        body = "\n\n".join(parts)
        return NormalizedArtifact(
            source_kind=SourceKind.OFFICE,
            title=Path(relative_path).name,
            relative_path=relative_path,
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            parser="pptx",
            parser_version=PARSER_VERSION,
            content_hash=sha256_text(body),
            content=ContentRef(text=body),
            structured_metadata={"slides": slides_meta},
            provenance={"relative_path": relative_path, "slides": slides_meta},
            outcome=MemberOutcome.SUCCESS,
        )

import json
import socket
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from platform_db import PlatformDatabase
from platform_services import KnowledgeService, PluginManager
from trading_service import PaperTradingService


class PlatformFeatureTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.db = PlatformDatabase(str(self.root / "hades.db"))
        self.db.initialize()
        self.knowledge = KnowledgeService(self.db, self.root / "data")

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_knowledge_is_persistent_and_searchable(self) -> None:
        source = self.knowledge.ingest_text(
            title="HADES architectuur",
            text="HADES gebruikt specialistische agents en een duurzame Knowledge Library voor projectkennis.",
            source_type="test",
            uri="test:hades",
        )
        self.assertGreater(source["chunks"], 0)
        matches = self.db.search_knowledge("specialistische agents", 5)
        self.assertTrue(matches)
        self.assertEqual(matches[0]["source_id"], source["id"])
        reopened = PlatformDatabase(str(self.root / "hades.db"))
        reopened.initialize()
        self.assertTrue(reopened.search_knowledge("Knowledge Library", 5))

    def test_folder_ingestion_skips_unchanged_files(self) -> None:
        docs = self.root / "docs"
        docs.mkdir()
        (docs / "one.md").write_text("# Een\nLokale HADES kennisbron over retrieval.", encoding="utf-8")
        first = self.knowledge.ingest_folder(docs, "Docs")
        second = self.knowledge.ingest_folder(docs, "Docs")
        self.assertEqual(first["ready"], 1)
        self.assertEqual(second["unchanged"], 1)

    def test_file_lifecycle_detail_rescan_and_delete(self) -> None:
        docs = self.root / "lifecycle-docs"
        docs.mkdir()
        note = docs / "note.md"
        note.write_text("# Lifecycle\nHADES bestanden indexeren en opnieuw scannen.", encoding="utf-8")
        uploaded = self.knowledge.store_upload("upload-note.md", b"# Upload\nGeuploade kennis voor de Bestanden pagina.")
        upload_result = self.knowledge.ingest_file(uploaded, None)
        self.assertEqual(upload_result["file"]["status"], "ready")
        self.assertIsNone(upload_result["file"]["workspace_id"])

        folder = self.knowledge.ingest_folder(docs, "Lifecycle")
        workspace_id = folder["workspace"]["id"]
        indexed = self.db.list_indexed_files(workspace_id)
        self.assertEqual(len(indexed), 1)
        file_id = indexed[0]["id"]

        detail = self.knowledge.file_detail(file_id)
        self.assertIsNotNone(detail)
        assert detail is not None
        self.assertGreaterEqual(len(detail["chunks"]), 1)
        self.assertEqual(detail["file"]["status"], "ready")

        uploads_only = self.db.list_indexed_files("uploads")
        self.assertTrue(any(item["id"] == upload_result["file"]["id"] for item in uploads_only))

        note.write_text("# Lifecycle\nHADES bestanden indexeren en opnieuw scannen. Extra inhoud.", encoding="utf-8")
        rescanned = self.knowledge.rescan_workspace(workspace_id)
        self.assertGreaterEqual(rescanned["ready"], 1)

        removed_file = self.knowledge.remove_indexed_file(upload_result["file"]["id"])
        self.assertEqual(removed_file["id"], upload_result["file"]["id"])
        self.assertIsNone(self.db.get_indexed_file(upload_result["file"]["id"]))
        self.assertFalse(self.db.search_knowledge("Geuploade kennis", 5))

        removed_workspace = self.knowledge.remove_workspace(workspace_id)
        self.assertEqual(removed_workspace["id"], workspace_id)
        self.assertEqual(removed_workspace["removed_files"], 1)
        self.assertIsNone(self.db.get_workspace(workspace_id))
        self.assertFalse(self.db.list_indexed_files(workspace_id))

    def test_office_documents_are_extracted_and_chunked(self) -> None:
        from openpyxl import Workbook
        from pptx import Presentation

        workbook_path = self.root / "sheet.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Onderwerp", "Waarde"])
        sheet.append(["HADES", "Knowledge Library"] )
        workbook.save(workbook_path)
        xlsx = self.knowledge.ingest_file(workbook_path)
        self.assertEqual(xlsx["file"]["status"], "ready")
        self.assertTrue(self.db.search_knowledge("Knowledge Library", 5))

        pptx_path = self.root / "slides.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "HADES Research"
        slide.placeholders[1].text = "Expert research bewaart bewijs in chunks."
        presentation.save(pptx_path)
        pptx = self.knowledge.ingest_file(pptx_path)
        self.assertEqual(pptx["file"]["status"], "ready")
        self.assertTrue(self.db.search_knowledge("bewijs chunks", 5))

    def test_plugin_manifest_pack_and_invoke(self) -> None:
        source = self.root / "echo-plugin"
        source.mkdir()
        (source / "echo.py").write_text(
            "import sys\nprint('PLUGIN_OK ' + ' '.join(sys.argv[1:]))\n",
            encoding="utf-8",
        )
        manifest = {
            "format": 1,
            "id": "echo-test",
            "name": "Echo Test",
            "version": "1.0.0",
            "runtime_type": "python",
            "entrypoint": "echo.py",
            "permissions": ["subprocess"],
            "tools": [
                {
                    "name": "echo",
                    "description": "Echo testwoorden via een lokale subprocess-tool.",
                    "input_schema": {"type": "object", "properties": {"args": {"type": "array", "items": {"type": "string"}}}},
                    "command": "{python} echo.py {args}",
                }
            ],
        }
        (source / "hades-plugin.json").write_text(json.dumps(manifest), encoding="utf-8")
        manager = PluginManager(self.db, self.root / "data")
        converted = manager.import_local_folder(source, install_dependencies=False)
        self.assertEqual(converted["plugin"]["status"], "ready")
        self.assertFalse(converted["plugin"]["enabled"])
        self.assertTrue(Path(converted["package_path"]).is_file())
        self.assertEqual(Path(converted["package_path"]).suffix, ".HadesPlugin")
        self.db.set_plugin_state(converted["plugin"]["id"], enabled=True)
        output = manager.invoke(converted["plugin"]["id"], "echo", {"args": ["hallo", "HADES"]})
        self.assertEqual(output["status"], "completed")
        self.assertIn("PLUGIN_OK hallo HADES", output["output"])
        self.assertEqual(self.db.get_plugin(converted["plugin"]["id"])["health"], "operational")
        exported = manager.export_package(converted["plugin"]["id"])
        self.assertTrue(exported.is_file())
        roundtrip = manager.import_zip(Path(converted["package_path"]), install_dependencies=False)
        self.assertEqual(roundtrip["plugin"]["status"], "ready")
        self.assertFalse(roundtrip["plugin"]["enabled"])
        self.assertEqual(Path(roundtrip["plugin"]["local_path"]).name, "source")
        self.db.set_plugin_state(roundtrip["plugin"]["id"], enabled=True)
        output2 = manager.invoke(roundtrip["plugin"]["id"], "echo", {"args": ["roundtrip"]})
        self.assertIn("PLUGIN_OK roundtrip", output2["output"])
        runtime_source = Path(roundtrip["plugin"]["local_path"])
        self.assertTrue(manager.uninstall(roundtrip["plugin"]["id"]))
        self.assertIsNone(self.db.get_plugin(roundtrip["plugin"]["id"]))
        self.assertFalse(runtime_source.exists())

    def test_failed_service_healthcheck_never_reports_started(self) -> None:
        with socket.socket() as probe:
            probe.bind(("127.0.0.1", 0))
            unused_port = probe.getsockname()[1]
        source = self.root / "unhealthy-service"
        source.mkdir()
        (source / "service.py").write_text("import time\ntime.sleep(30)\n", encoding="utf-8")
        manifest = {
            "format": 1,
            "id": "unhealthy-service",
            "name": "Unhealthy Service",
            "version": "1.0.0",
            "runtime_type": "python",
            "permissions": ["subprocess"],
            "healthcheck": {"type": "tcp", "host": "127.0.0.1", "port": unused_port, "timeout_seconds": 0.25, "interval_seconds": 0.05},
            "tools": [{"name": "start", "action": "start", "mode": "service", "command": ["{python}", "service.py"], "input_schema": {"type": "object", "properties": {}}}],
        }
        (source / "hades-plugin.json").write_text(json.dumps(manifest), encoding="utf-8")
        manager = PluginManager(self.db, self.root / "data")
        converted = manager.import_local_folder(source, install_dependencies=False)
        self.db.set_plugin_state(converted["plugin"]["id"], enabled=True)
        result = manager.invoke(converted["plugin"]["id"], "start", {}, timeout=1, invocation_type="manual", approved_by_user=True)
        self.assertEqual(result["status"], "failed")
        self.assertIn("niet bereikbaar", result["error"])
        self.assertEqual(self.db.get_plugin(converted["plugin"]["id"])["health"], "unhealthy")
        self.assertNotIn(converted["plugin"]["id"], manager._processes)
        persisted = self.db.tool_calls(converted["plugin"]["id"])
        self.assertEqual(persisted[0]["status"], "failed")
        self.assertIsNotNone(persisted[0]["duration_ms"])

    def test_http_healthcheck_does_not_follow_redirects(self) -> None:
        from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
        import threading

        class Handler(BaseHTTPRequestHandler):
            def do_GET(self):  # noqa: N802
                self.send_response(302)
                self.send_header("Location", "http://169.254.169.254/latest/meta-data/")
                self.end_headers()

            def log_message(self, *_args):
                return

        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            manager = PluginManager(self.db, self.root / "data")
            plugin = {"id": "health-redirect", "local_path": str(self.root)}
            result = manager._healthcheck(
                plugin,
                {"type": "http", "url": f"http://127.0.0.1:{server.server_address[1]}/"},
                timeout=2.0,
            )
            self.assertFalse(result["healthy"], result)
            self.assertIn("redirect", (result.get("detail") or "").lower())
        finally:
            server.shutdown()
            server.server_close()

    def test_service_lifecycle_build_start_health_status_logs_and_stop(self) -> None:
        with socket.socket() as probe:
            probe.bind(("127.0.0.1", 0))
            port = probe.getsockname()[1]
        source = self.root / "healthy-service"
        source.mkdir()
        (source / "service.py").write_text(
            "import socket, sys\n"
            "server = socket.socket()\n"
            "server.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)\n"
            "server.bind(('127.0.0.1', int(sys.argv[1])))\n"
            "server.listen()\n"
            "print('SERVICE_READY', flush=True)\n"
            "while True:\n"
            "    client, _ = server.accept()\n"
            "    client.close()\n",
            encoding="utf-8",
        )
        empty_schema = {"type": "object", "properties": {}}
        manifest = {
            "format": 1,
            "id": "healthy-service",
            "name": "Healthy Service",
            "version": "1.0.0",
            "runtime_type": "python",
            "permissions": ["subprocess"],
            "healthcheck": {"type": "tcp", "host": "127.0.0.1", "port": port, "timeout_seconds": 2, "interval_seconds": 0.05},
            "tools": [
                {"name": "build", "action": "build", "command": ["{python}", "-c", "print('BUILD_OK')"], "input_schema": empty_schema},
                {"name": "start", "action": "start", "mode": "service", "command": ["{python}", "service.py", "{port}"], "input_schema": {"type": "object", "required": ["port"], "properties": {"port": {"type": "integer"}}}},
                {"name": "health", "action": "health", "command": "", "input_schema": empty_schema},
                {"name": "status", "action": "status", "command": "", "input_schema": empty_schema},
                {"name": "logs", "action": "logs", "command": "", "input_schema": empty_schema},
                {"name": "stop", "action": "stop", "command": "", "input_schema": empty_schema},
            ],
        }
        (source / "hades-plugin.json").write_text(json.dumps(manifest), encoding="utf-8")
        manager = PluginManager(self.db, self.root / "data")
        plugin = manager.import_local_folder(source, install_dependencies=False)["plugin"]
        self.db.set_plugin_state(plugin["id"], enabled=True)
        self.assertEqual(manager.invoke(plugin["id"], "build", {})["status"], "completed")
        started = manager.invoke(plugin["id"], "start", {"port": port}, timeout=3)
        self.assertEqual(started["status"], "completed", started.get("error"))
        self.assertEqual(self.db.get_plugin(plugin["id"])["health"], "healthy")
        self.assertEqual(manager.invoke(plugin["id"], "health", {})["status"], "completed")
        self.assertEqual(manager.invoke(plugin["id"], "status", {})["status"], "completed")
        logs = manager.invoke(plugin["id"], "logs", {})
        self.assertIn("SERVICE_READY", logs["stdout"])
        stopped = manager.invoke(plugin["id"], "stop", {}, timeout=3)
        self.assertEqual(stopped["status"], "completed", stopped.get("error"))
        self.assertEqual(self.db.get_plugin(plugin["id"])["health"], "stopped")

    def test_dependency_failure_is_persistent_and_plugin_stays_disabled(self) -> None:
        source = self.root / "broken-dependencies"
        source.mkdir()
        (source / "package.json").write_text('{"name":"broken-dependencies","scripts":{"build":"node build.js"}}', encoding="utf-8")
        manifest = {
            "format": 1,
            "id": "broken-dependencies",
            "name": "Broken Dependencies",
            "version": "1.0.0",
            "runtime_type": "node",
            "permissions": ["subprocess"],
            "tools": [{"name": "build", "action": "build", "command": ["{npm}", "run", "build"], "input_schema": {"type": "object", "properties": {}}}],
        }
        (source / "hades-plugin.json").write_text(json.dumps(manifest), encoding="utf-8")
        failed = {
            "command": ["npm", "install"],
            "stdout": "",
            "stderr": "dependency exploded",
            "exit_code": 1,
            "error": "dependency exploded",
            "timed_out": False,
            "stalled": False,
            "duration_ms": 12,
        }
        manager = PluginManager(self.db, self.root / "data")
        with patch("platform_services.shutil.which", return_value="npm"), patch(
            "platform_services.DependencyCommandRunner.run", return_value=failed
        ):
            converted = manager.import_local_folder(source, install_dependencies=True)
        self.assertEqual(converted["plugin"]["status"], "needs_review")
        self.assertFalse(converted["plugin"]["enabled"])
        self.assertIn("dependency exploded", converted["plugin"]["last_error"])
        dependency_call = next(call for call in self.db.tool_calls(converted["plugin"]["id"]) if call["tool_name"] == "__dependencies__")
        self.assertEqual(dependency_call["status"], "failed")
        self.assertIn("dependency exploded", dependency_call["stderr"])

    def test_tool_arguments_are_never_interpreted_by_a_shell(self) -> None:
        source = self.root / "argv-safety"
        source.mkdir()
        marker = self.root / "injected.txt"
        # Print argv as raw lines (not repr) so Windows paths with backslashes
        # remain assertable, and report argc so shell-splitting would fail loudly.
        (source / "argv.py").write_text(
            "import sys\nprint('ARGC=' + str(len(sys.argv) - 1))\nfor arg in sys.argv[1:]:\n    print(arg)\n",
            encoding="utf-8",
        )
        manifest = {
            "format": 1,
            "id": "argv-safety",
            "name": "Argv Safety",
            "version": "1.0.0",
            "runtime_type": "python",
            "permissions": ["subprocess"],
            "tools": [{"name": "run", "command": ["{python}", "argv.py", "{args}"], "input_schema": {"type": "object", "properties": {"args": {"type": "array", "items": {"type": "string"}}}}}],
        }
        (source / "hades-plugin.json").write_text(json.dumps(manifest), encoding="utf-8")
        manager = PluginManager(self.db, self.root / "data")
        converted = manager.import_local_folder(source, install_dependencies=False)
        self.db.set_plugin_state(converted["plugin"]["id"], enabled=True)
        payload = f"safe; touch {marker}"
        result = manager.invoke(converted["plugin"]["id"], "run", {"args": [payload]}, approved_by_user=True)
        self.assertEqual(result["status"], "completed")
        self.assertIn("ARGC=1", result["stdout"])
        self.assertIn(payload, result["stdout"])
        self.assertFalse(marker.exists())

    def test_plugin_zip_blocks_path_traversal(self) -> None:
        import zipfile
        archive = self.root / "bad.zip"
        with zipfile.ZipFile(archive, "w") as zf:
            zf.writestr("../escape.txt", "no")
        target = self.root / "extract"
        target.mkdir()
        with self.assertRaises(ValueError):
            PluginManager.safe_extract_zip(archive, target)
        self.assertFalse((self.root / "escape.txt").exists())

    def test_plugin_manifest_requires_identity_and_tool_inputs_are_validated(self) -> None:
        source = self.root / "invalid-manifest"
        source.mkdir()
        (source / "hades-plugin.json").write_text(json.dumps({
            "format": 1, "name": "Invalid", "version": "1.0.0", "runtime_type": "python", "tools": [],
        }), encoding="utf-8")
        manager = PluginManager(self.db, self.root / "data")
        with self.assertRaisesRegex(ValueError, "verplichte velden"):
            manager.import_local_folder(source, install_dependencies=False)

        valid = self.root / "validated-plugin"
        valid.mkdir()
        (valid / "echo.py").write_text("import sys\nprint(sys.argv[1])\n", encoding="utf-8")
        (valid / "hades-plugin.json").write_text(json.dumps({
            "format": 1, "id": "validated-plugin", "name": "Validated", "version": "1.0.0",
            "runtime_type": "python", "permissions": ["subprocess"],
            "tools": [{"name": "echo", "command": ["{python}", "echo.py", "{mode}"],
                       "input_schema": {"type": "object", "required": ["mode"], "properties": {
                           "mode": {"type": "string", "enum": ["safe"]},
                       }, "additionalProperties": False}}],
        }), encoding="utf-8")
        plugin = manager.import_local_folder(valid, install_dependencies=False)["plugin"]
        self.db.set_plugin_state(plugin["id"], enabled=True)
        invalid = manager.invoke(plugin["id"], "echo", {"mode": "unsafe"}, approved_by_user=True)
        self.assertEqual(invalid["status"], "failed")
        self.assertIn("toegestane waarden", invalid["error"])
        unknown = manager.invoke(plugin["id"], "echo", {"mode": "safe", "extra": True}, approved_by_user=True)
        self.assertEqual(unknown["status"], "failed")
        self.assertIn("Onbekende toolvelden", unknown["error"])

    def test_tool_call_inputs_redact_secret_fields(self) -> None:
        call_id = self.db.create_tool_call(
            "plugin", "tool", {"query": "visible", "api_key": "top-secret", "nested": {"token": "also-secret"}},
            metadata={"authorization": "Bearer secret"},
        )
        call = self.db.get_tool_call(call_id)
        self.assertEqual(call["input"]["query"], "visible")
        self.assertEqual(call["input"]["api_key"], "[REDACTED]")
        self.assertEqual(call["input"]["nested"]["token"], "[REDACTED]")
        self.assertEqual(call["metadata"]["authorization"], "[REDACTED]")

    def test_status_command_without_healthcheck_is_operational_not_healthy(self) -> None:
        source = self.root / "status-command"
        source.mkdir()
        (source / "hades-plugin.json").write_text(json.dumps({
            "format": 1, "id": "status-command", "name": "Status Command", "version": "1.0.0",
            "runtime_type": "python", "permissions": ["subprocess"],
            "tools": [{"name": "status", "action": "status", "command": ["{python}", "-c", "print('OK')"],
                       "input_schema": {"type": "object", "properties": {}}}],
        }), encoding="utf-8")
        manager = PluginManager(self.db, self.root / "data")
        plugin = manager.import_local_folder(source, install_dependencies=False)["plugin"]
        self.db.set_plugin_state(plugin["id"], enabled=True)
        result = manager.invoke(plugin["id"], "status", {}, approved_by_user=True)
        self.assertEqual(result["status"], "completed")
        self.assertEqual(self.db.get_plugin(plugin["id"])["health"], "operational")

    def test_restart_reconciliation_clears_unverifiable_healthy_state(self) -> None:
        source = self.root / "stale-service"
        source.mkdir()
        (source / "hades-plugin.json").write_text(json.dumps({
            "format": 1, "id": "stale-service", "name": "Stale Service", "version": "1.0.0",
            "runtime_type": "python", "permissions": ["subprocess"],
            "tools": [{"name": "run", "command": ["{python}", "-c", "print('OK')"],
                       "input_schema": {"type": "object", "properties": {}}}],
        }), encoding="utf-8")
        manager = PluginManager(self.db, self.root / "data")
        plugin = manager.import_local_folder(source, install_dependencies=False)["plugin"]
        self.db.set_plugin_state(plugin["id"], health="healthy")
        PluginManager(self.db, self.root / "data").reconcile_services()
        restored = self.db.get_plugin(plugin["id"])
        self.assertEqual(restored["health"], "needs_attention")
        self.assertIn("na herstart", restored["last_error"])

    def test_gods_eye_view_manifest_imports_as_ready_service_plugin(self) -> None:
        """Shipped GEV HADES manifest must convert with lifecycle tools + healthcheck."""
        repo_root = Path(__file__).resolve().parents[2]
        manifest_path = repo_root / "plugins" / "gods-eye-view" / "hades-plugin.json"
        self.assertTrue(manifest_path.is_file(), "plugins/gods-eye-view/hades-plugin.json ontbreekt")
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        self.assertEqual(manifest.get("format"), 1)
        self.assertEqual(manifest.get("id"), "gods-eye-view")
        self.assertFalse(manifest.get("autonomous", True))
        self.assertEqual(manifest.get("healthcheck", {}).get("type"), "http")
        self.assertIn("4173", str(manifest.get("healthcheck", {}).get("url", "")))

        source = self.root / "gods-eye-view-fixture"
        source.mkdir()
        (source / "hades-plugin.json").write_text(json.dumps(manifest), encoding="utf-8")
        (source / "package.json").write_text(
            json.dumps(
                {
                    "name": "gods-eye-view",
                    "version": "0.1.1",
                    "private": True,
                    "scripts": {"dev": "node -e \"process.stdout.write('dev')\"", "doctor": "node -e \"process.stdout.write('doctor')\""},
                }
            ),
            encoding="utf-8",
        )
        manager = PluginManager(self.db, self.root / "data")
        converted = manager.import_local_folder(source, install_dependencies=False)
        plugin = converted["plugin"]
        self.assertEqual(plugin["status"], "ready")
        self.assertFalse(plugin["enabled"])
        self.assertEqual(plugin["runtime_type"], "node")
        tool_names = {tool["name"] for tool in converted["tools"]}
        self.assertTrue({"doctor", "start", "health", "status", "logs", "stop"} <= tool_names)
        start = next(tool for tool in converted["tools"] if tool["name"] == "start")
        self.assertEqual(start["metadata"].get("action"), "start")
        self.assertEqual(start["metadata"].get("mode"), "service")
        self.assertEqual(start["metadata"].get("env", {}).get("PORT"), "4173")

    def test_uploaded_folder_is_packed_into_hadesplugin(self) -> None:
        manager = PluginManager(self.db, self.root / "data")
        self.assertIsNone(manager.resolve_upload_relative("plugin/node_modules/skip.js"))
        with self.assertRaises(ValueError):
            manager.resolve_upload_relative("../evil.py")
        with self.assertRaises(ValueError):
            manager.resolve_upload_relative("C:/Windows/system32/evil.py")
        payload = [
            ("folder-echo/hades-plugin.json", json.dumps({
                "format": 1,
                "id": "folder-echo",
                "name": "Folder Echo",
                "version": "1.0.0",
                "runtime_type": "python",
                "permissions": ["subprocess"],
                "tools": [{"name": "echo", "command": ["{python}", "echo.py", "{args}"], "input_schema": {"type": "object", "properties": {"args": {"type": "array", "items": {"type": "string"}}}}}],
            }).encode("utf-8")),
            ("folder-echo/echo.py", b"import sys\nprint('FOLDER ' + ' '.join(sys.argv[1:]))\n"),
            ("folder-echo/node_modules/left-behind.js", b"should-skip"),
        ]
        converted = manager.import_uploaded_folder(payload, install_dependencies=False)
        self.assertEqual(converted["plugin"]["status"], "ready")
        self.assertFalse(converted["plugin"]["enabled"])
        self.assertTrue(Path(converted["package_path"]).is_file())
        self.assertEqual(Path(converted["package_path"]).suffix, ".HadesPlugin")
        self.assertFalse((Path(converted["plugin"]["local_path"]) / "node_modules").exists())
        self.db.set_plugin_state(converted["plugin"]["id"], enabled=True)
        output = manager.invoke(converted["plugin"]["id"], "echo", {"args": ["packed"]})
        self.assertEqual(output["status"], "completed")
        self.assertIn("FOLDER packed", output["output"])

    def test_paper_trading_is_transactional_and_exactly_once(self) -> None:
        paper = PaperTradingService(self.db)
        paper.set_enabled(True)
        opened = paper.buy("BTC/USDT", 0.1, 10_000)
        self.assertEqual(opened["price_source"], "operator_input")
        self.assertEqual(opened["simulation_layer"], "legacy_paper_desk")
        after_open = opened["state"]
        self.assertAlmostEqual(after_open["wallets"][0]["balance"], 9_000)
        position_id = opened["position_id"]
        closed = paper.close(position_id, 11_000)
        self.assertEqual(closed["price_source"], "operator_input")
        self.assertAlmostEqual(closed["state"]["wallets"][0]["balance"], 10_100)
        self.assertAlmostEqual(closed["realized_pnl"], 100)
        with self.assertRaises(RuntimeError):
            paper.close(position_id, 11_000)
        self.assertAlmostEqual(paper.state()["wallets"][0]["balance"], 10_100)

    def test_paper_kill_switch_blocks_new_orders(self) -> None:
        paper = PaperTradingService(self.db)
        paper.set_enabled(True)
        paper.set_kill_switch(True)
        with self.assertRaises(RuntimeError):
            paper.buy("ETH/USDT", 1, 100)

    def test_trading_bot_discovers_strategies_and_builds_knowledge(self) -> None:
        from trading_service import TradingBotService

        paper = PaperTradingService(self.db)
        bot = TradingBotService(self.db, paper, self.knowledge)
        seeded = bot.seed_synthetic(symbol="BTC/USDT", bars=180, seed=7)
        self.assertEqual(seeded["bars"], 180)
        bars = bot.get_bars("BTC/USDT", "1h", 500)
        self.assertGreaterEqual(len(bars), 180)
        discoveries = bot.discover_strategies(bars, top_n=3)
        self.assertEqual(len(discoveries), 3)
        self.assertIn(discoveries[0]["kind"], TradingBotService.STRATEGY_KINDS)
        saved = bot.persist_discovered(discoveries, symbol="BTC/USDT", timeframe="1h")
        self.assertEqual(len(saved), 3)
        self.assertEqual(saved[0]["status"], "active")
        learning = bot.ingest_learning(
            title="Trading strategie-ontdekking · BTC/USDT",
            text="# Test learning\n\nDeterministische PAPER-strategiekennis voor HADES.",
            uri="trading://strategy-run/test",
            metadata={"kind": "discover"},
        )
        self.assertIsNotNone(learning)
        assert learning is not None
        matches = self.db.search_knowledge("PAPER-strategiekennis", 5)
        self.assertTrue(matches)
        self.assertEqual(matches[0]["source_id"], learning["id"])
        bot.set_bot_settings(strategy_id=saved[0]["id"], symbol="BTC/USDT", enabled=False)
        paper.set_enabled(True)
        step = bot.execute_paper_bot_step(saved[0], bars, position_fraction=0.1)
        self.assertIn(step["signal"], (0, 1))
        self.assertTrue(step["actions"])
        paper.set_enabled(False)
        with self.assertRaisesRegex(RuntimeError, "paper_trading_disabled"):
            bot.execute_paper_bot_step(saved[0], bars, position_fraction=0.1)
        dash = bot.dashboard()
        self.assertTrue(dash["strategies"])
        self.assertTrue(dash["bars"])
        self.assertEqual(dash["bot"]["strategy_id"], saved[0]["id"])

    def test_trading_csv_import_and_backtest(self) -> None:
        from trading_service import TradingBotService

        paper = PaperTradingService(self.db)
        bot = TradingBotService(self.db, paper, self.knowledge)
        lines = ["ts,open,high,low,close,volume"]
        price = 100.0
        for index in range(80):
            day = 1 + (index // 24)
            hour = index % 24
            nxt = price * (1.01 if index % 7 else 0.99)
            high = max(price, nxt) * 1.002
            low = min(price, nxt) * 0.998
            lines.append(
                f"2024-01-{day:02d}T{hour:02d}:00:00+00:00,{price:.4f},{high:.4f},{low:.4f},{nxt:.4f},10"
            )
            price = nxt
        summary = bot.import_csv("\n".join(lines), "ETH/USDT", "1h")
        self.assertEqual(summary["bars"], 80)
        bars = bot.get_bars("ETH/USDT", "1h", 200)
        metrics = bot.backtest(bars, kind="sma_crossover", params={"fast": 5, "slow": 20})
        self.assertIn("score", metrics)
        self.assertEqual(metrics["bars"], 80)
        self.assertEqual(metrics["simulation_layer"], "legacy_paper_desk")
        self.assertEqual(metrics["research_class"], "educational_in_sample")
        self.assertFalse(metrics["validated_strategy"])
        self.assertIn("same_bar_signal_same_bar_fill_lookahead", metrics["limitations"])
        self.assertIn("Trading Lab", metrics["disclaimer"])
        discoveries = bot.discover_strategies(bars, top_n=1)
        self.assertFalse(discoveries[0]["validated_strategy"])
        saved = bot.persist_discovered(discoveries, symbol="ETH/USDT", timeframe="1h")
        self.assertIn("[legacy/educational/in-sample]", saved[0]["notes"])
        self.assertIn("Geen gevalideerde strategie", saved[0]["notes"])


if __name__ == "__main__":
    unittest.main()
